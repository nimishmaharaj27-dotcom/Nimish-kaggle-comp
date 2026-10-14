"""Convert report.md (simple markdown subset) to a Word document with figures and tables."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

here = Path(__file__).parent
doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)

def runs(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"): par.add_run(part[2:-2]).bold = True
        elif part.startswith("*"): par.add_run(part[1:-1]).italic = True
        elif part.startswith("`"): par.add_run(part[1:-1]).font.name = "Consolas"
        else: par.add_run(part)

lines = (here / "report.md").read_text(encoding="utf-8").split("\n"); i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("# "): doc.add_heading(ln[2:], 0)
    elif ln.startswith("## "): doc.add_heading(ln[3:], 1)
    elif ln.startswith("!["):
        m = re.match(r"!\[(.*?)\]\((.*?)\)", ln)
        doc.add_picture(str((here / m.group(2)).resolve()), width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = doc.add_paragraph(m.group(1)); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic = True
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[-| ]+\|$", lines[i]): rows.append([c.strip() for c in lines[i].strip("|").split("|")])
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c); cell.text = ""; runs(cell.paragraphs[0], val)
                for rn in cell.paragraphs[0].runs: rn.font.size = Pt(9)
        continue
    elif re.match(r"^\d+\. ", ln): runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", ln))
    elif ln.strip(): runs(doc.add_paragraph(), ln)
    i += 1
doc.save(here / "ENGE707_A3_Report.docx"); print("saved")
